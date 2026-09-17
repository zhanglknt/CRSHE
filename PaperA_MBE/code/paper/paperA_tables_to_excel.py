#!/usr/bin/env python3
"""
Paper A (MBE) - Build standalone Excel table files for submission.
Main tables 1-4 -> results/paper/Main_Tables_v9.xlsx
Supplementary text tables (S2-1, S3-1, S4-1/2/3) -> results/paper/supplementary/Supplementary_Text_Tables_v9.xlsx
Note: Table 2 'Unclassified / Current' corrected 68.5% -> 68.4% (3,404/4,974 = 68.44%,
consistent with Table 1 and the round-4 R3 numeric-fidelity fix).
"""
import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

if os.path.exists("/mnt/d"):
    BASE = "/mnt/d/人类正选择基因项目"
else:
    BASE = "D:/人类正选择基因项目"

TITLE_FONT = Font(bold=True, size=12)
HEADER_FONT = Font(bold=True)
NOTE_FONT = Font(italic=True, size=10)


# ------------------------------------------------------------------
# Table S3-1 data (P0-D corrected Selectome cross-reference):
# built from the raw Selectome v6 NHX parse (primate ancestral branches)
# instead of the invalid summary_v3 family->gene mapping. The nine genes
# are those with Selectome primate-ancestral support AND BUSTED FDR < 0.05
# in the 4,974-gene universe.
# ------------------------------------------------------------------
strip = lambda s: str(s).split(".")[0]
_gc = pd.read_csv(f"{BASE}/results/phase7_gene_vs_regulation/phase7g_classification_v7/"
                  "gene_classification_v7.csv", low_memory=False)
_gc["ensg"] = _gc["gene_id"].map(strip)
_pps = pd.read_csv(f"{BASE}/data/selectome/selectome_primate_positive_selection.tsv", sep="\t")
_pps["ensg"] = _pps["gene_id"].map(strip)
_hps = pd.read_csv(f"{BASE}/data/selectome/selectome_human_positive_selection.tsv", sep="\t")
_hps["ensg"] = _hps["gene_id"].map(strip)
_desc = dict(zip(_hps["ensg"], _hps["gene_description"].astype(str)))
_gmap = pd.read_csv(f"{BASE}/data/gene_id_to_symbol_gencode47.csv")
_id2sym = {strip(g): s for g, s in zip(_gmap["gene_id"], _gmap["gene_symbol"])}
_gc["sym"] = _gc["gene_symbol"].where(
    _gc["gene_symbol"].notna() & (_gc["gene_symbol"].astype(str).str.strip() != ""),
    _gc["ensg"].map(_id2sym)).astype(str)
_sel = _gc[_gc["ensg"].isin(set(_pps["ensg"])) & _gc["bh_fdr_sig"].astype(bool)].copy()
_sel = _sel.sort_values("sym")
FUNC = {  # concise functional annotation (HGNC descriptions)
    "CCDC122": "Coiled-coil domain containing 122",
    "CCHCR1": "Coiled-coil alpha-helical rod protein 1",
    "CD14": "CD14 molecule, monocyte/innate immune receptor",
    "MRPS5": "Mitochondrial ribosomal protein S5",
    "NUB1": "Negative regulator of ubiquitin-like proteins 1",
    "PTER": "Phosphotriesterase-related protein",
    "RAD17": "RAD17 checkpoint clamp loader, DNA damage response",
    "SLC38A9": "Solute carrier 38A9, lysosomal arginine sensor (mTORC1)",
    "TMEM171": "Transmembrane protein 171",
}
tier1_rows = []
for _, r in _sel.iterrows():
    pp = _pps[_pps["ensg"] == r["ensg"]]
    branch = "; ".join(str(x) for x in pp["primate_pos_details"]) if len(pp) else ""
    tier1_rows.append([
        r["sym"],
        float(r["busted_p"]), float(r["busted_lrt"]),
        branch, float(r["relax_K"]) if pd.notna(r["relax_K"]) else None,
        FUNC.get(r["sym"].upper(), _desc.get(r["ensg"], ""))])
assert len(tier1_rows) == 9, f"expected 9 Tier-1 genes, got {len(tier1_rows)}"


def write_sheet(ws, title, headers, rows, notes=None):
    ws.cell(row=1, column=1, value=title).font = TITLE_FONT
    hr = 3
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=hr, column=j, value=h)
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center")
    for i, row in enumerate(rows, start=hr + 1):
        for j, v in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=v)
    r = hr + len(rows) + 2
    if notes:
        for note in notes:
            ws.cell(row=r, column=1, value=note).font = NOTE_FONT
            r += 1
    # column widths
    for j in range(1, len(headers) + 1):
        vals = [str(headers[j - 1])] + [str(row[j - 1]) for row in rows]
        w = min(max(len(v) for v in vals) + 3, 60)
        ws.column_dimensions[get_column_letter(j)].width = w


# ================= Main Tables 1-4 =================
wb = Workbook()

write_sheet(
    wb.active, "Table 1. Classification summary (4,974 genes).",
    ["Class", "n", "%"],
    [["Gene-driven", 1214, "24.4%"],
     ["Gene-driven (relaxed)", 52, "1.0%"],
     ["Regulation-driven", 293, "5.9%"],
     ["Dual-driven", 11, "0.2%"],
     ["Unclassified", 3404, "68.4%"]])
wb.active.title = "Table 1"

write_sheet(
    wb.create_sheet("Table 2"),
    "Table 2. Effect of comparability correction (previous -> current framework).",
    ["Class", "Previous", "Current", "Change"],
    [["Gene-driven", "1,620 (32.6%)", "1,214 (24.4%)", "-406 (-25.1%)"],
     ["Gene-driven (relaxed)", "-", "52 (1.0%)", "new category"],
     ["Regulation-driven", "125 (2.5%)", "293 (5.9%)", "+168 (+134.4%)"],
     ["Dual-driven", "19 (0.4%)", "11 (0.2%)", "-8 (-42.1%)"],
     ["Gene-driven (dual)", "57 (1.1%)", "-", "class removed"],
     ["Unclassified", "3,153 (63.4%)", "3,404 (68.4%)", "+251 (+8.0%)"]])

write_sheet(
    wb.create_sheet("Table 3"),
    "Table 3. Complete leave-one-out matrix (GD excludes gene-driven [relaxed]; Fisher exact, 95% CI).",
    ["Variant", "GD", "RD", "HAR OR (95% CI)", "hCONDEL OR (95% CI)"],
    [["Full", 1214, 293, "4.15 (3.15-5.46)***", "2.94 (1.92-4.51)***"],
     ["LOO-caMPRA", 1020, 800, "1.61 (1.28-2.03)***", "2.27 (1.64-3.15)***"],
     ["LOO-tau", 1339, 148, "6.39 (4.51-9.03)***", "3.91 (2.33-6.56)***"],
     ["LOO-nc", 1369, 110, "9.39 (6.37-13.83)***", "3.36 (1.81-6.24)***"],
     ["LOO-brain", 1099, 613, "2.78 (2.21-3.48)***", "2.27 (1.60-3.24)***"]],
    notes=["***p < 0.001 (all survive Holm correction across variants).",
           "Class membership is threshold-sensitive - RD ranges from 110 (LOO-nc) to 800 (LOO-caMPRA),",
           "and per-gene median stability is 0.833 for RD versus 1.000 for GD -",
           "but the enrichment signal is robust to the removal of any single component.",
           "95% CIs by normal approximation on log(OR). The companion manuscript reports the",
           "Fisher conditional-MLE interval for the full-variant hCONDEL test (OR 2.94, 95% CI 1.85-4.55);",
           "both intervals lead to the same conclusion."])

write_sheet(
    wb.create_sheet("Table 4"),
    "Table 4. Covariate sensitivity of the CDS-length residualization.",
    ["Residualization", "Agreement", "GD", "GD (relaxed)", "RD", "Dual"],
    [["log10 CDS length (baseline)", "99.96%*", 1216, 52, 293, 11],
     ["+ GC3", "98.79%", 1213, 52, 304, 11],
     ["+ log10 median expression", "97.87%", 1214, 50, 275, 11],
     ["+ GC3 + expression", "97.55%", 1218, 52, 276, 12]],
    notes=["*Baseline row replicates the original pipeline by independent reimplementation (script provided in the code repository);",
           "the residual 0.04% reflects percentile-convention edge cases.",
           "The three covariate rows are computed by the same independent implementation rather than the original pipeline."])

out1 = f"{BASE}/results/paper/Main_Tables_v9.xlsx"
wb.save(out1)
print("saved", out1)

# ================= Supplementary Text Tables =================
wb2 = Workbook()

write_sheet(
    wb2.active, "Table S2-1. MEME site-level results",
    ["Gene", "Sequences", "Sites tested", "Sig. sites (p<0.05)", "Sig. sites (p<0.01)"],
    [["RBM20", 8, 525, 9, 0],
     ["C4orf48", 10, 530, 9, 2],
     ["CEP350", 9, 482, 7, 0],
     ["SH3TC2", 10, 363, 6, 1],
     ["SDK1", 10, 418, 4, 2],
     ["SH2D3C", 9, 343, 4, 2],
     ["STIP1", 9, 255, 1, 1],
     ["ZNF678", 9, 803, 0, 0],
     ["Total", "-", 3719, "40 (1.08%)", "8 (0.22%)"]])
wb2.active.title = "S2-1 MEME"

write_sheet(
    wb2.create_sheet("S3-1 Tier1"),
    "Table S3-1. Genes detected by both BUSTED (FDR < 0.05) and Selectome v6 "
    "(primate ancestral branches)",
    ["Gene", "BUSTED p", "BUSTED LRT", "Selectome branch (branch-site p)", "RELAX K", "Function"],
    tier1_rows,
    notes=["Selectome support: positive selection on primate ancestral branches in the Selectome v6",
           "database (NHX tree parse; branch-site p-values as reported by Selectome).",
           "RAD17 shows support on three nested primate branches (Homininae, Catarrhini, Simiiformes;",
           "p < 1e-200, saturation-level values). MRPS5 and PTER are classified neutral under v9;",
           "see manuscript Limitations for the 3-gene sensitivity of the classification to the",
           "Selectome flag. Site-level evidence is not part of the NHX extract (branch-level calls only)."])

write_sheet(
    wb2.create_sheet("S4-1 Top GD"),
    "Table S4-1. Top 10 gene-driven genes by GDS (RELAX K >= 50 = censored upper boundary)",
    ["Gene", "GDS", "BUSTED LRT", "RELAX K", "Function"],
    [["GLRX", 0.914, 64.9, 7.59, "Glutaredoxin, redox regulation"],
     ["E4F1", 0.886, 86.7, 1.64, "E4F transcription factor, cell cycle"],
     ["TAPT1", 0.854, 88.1, 1.00, "Transmembrane adaptor"],
     ["GNRHR", 0.846, 286.0, 10.56, "Gonadotropin-releasing hormone receptor"],
     ["RNF151", 0.844, 208.6, 6.78, "RING finger protein, spermatogenesis"],
     ["RPL35", 0.842, 193.6, 4.41, "Ribosomal protein L35"],
     ["YARS2", 0.842, 260.4, 13.86, "Tyrosyl-tRNA synthetase 2, mitochondrial"],
     ["ARF6", 0.840, 171.9, ">=50 (censored)", "ARF GTPase 6, vesicle trafficking"],
     ["AMBN", 0.840, 255.4, ">=50 (censored)", "Ameloblastin, tooth enamel"],
     ["TMEM248", 0.837, 203.8, 11.13, "Transmembrane protein 248"]])

write_sheet(
    wb2.create_sheet("S4-2 Top RD"),
    "Table S4-2. Top 10 regulation-driven genes by RDS",
    ["Gene", "RDS", "HARs", "caMPRA", "tau", "Function"],
    [["CLSTN2", 0.872, 1, "Yes", 0.959, "Calsyntenin 2, synaptic adhesion"],
     ["ESRRG", 0.730, 3, "Yes", 0.923, "Estrogen-related receptor gamma"],
     ["TBX5", 0.719, 1, "Yes", 0.939, "T-box transcription factor, heart/limb"],
     ["TMEFF2", 0.716, 1, "Yes", 0.950, "Tomoregulin, neural regulation"],
     ["GALNTL6", 0.711, 2, "Yes", 0.908, "GalNAc-transferase-like, glycosylation"],
     ["CDH10", 0.708, 1, "Yes", 0.966, "Cadherin 10, neural adhesion"],
     ["DAB1", 0.708, 6, "Yes", 0.909, "Reelin signaling, neural migration"],
     ["PRDM16", 0.693, 2, "Yes", 0.900, "PR domain, brown adipose tissue"],
     ["TYR", 0.691, 1, "Yes", 0.997, "Tyrosinase, melanin synthesis"],
     ["CLIC5", 0.690, 2, "Yes", 0.889, "Chloride intracellular channel 5"]])

write_sheet(
    wb2.create_sheet("S4-3 Dual"),
    "Table S4-3. Dual-driven genes (11 genes)",
    ["Gene", "BUSTED LRT", "GDS", "RDS", "RELAX K", "HARs", "caMPRA", "Function"],
    [["SLC5A7", 46.4, 0.48, 0.68, "0.00*", 2, "Yes", "Choline transporter, cholinergic neurons"],
     ["STRA8", 31.9, 0.43, 0.67, "0.00*", 1, "Yes", "Meiosis initiator"],
     ["USP46", 39.7, 0.40, 0.67, 0.23, 1, "Yes", "Ubiquitin-specific peptidase"],
     ["TCEA3", 43.8, 0.46, 0.62, "0.00*", 1, "Yes", "Transcription elongation factor A3"],
     ["VWA5B1", 44.1, 0.46, 0.62, 2.14, 1, "Yes", "von Willebrand factor A domain"],
     ["STIP1", 48.0, 0.38, 0.61, "0.00*", 1, "Yes", "Stress-induced phosphoprotein 1"],
     ["CNTN4", 41.9, 0.32, 0.65, 0.78, 7, "Yes", "Contactin 4, axon guidance"],
     ["CPNE4", 49.2, 0.37, 0.66, 0.31, 0, "No", "Copine 4, calcium-dependent membrane binding"],
     ["TMEM71", 21.5, 0.37, 0.53, 0.44, 1, "Yes", "Transmembrane protein 71"],
     ["VPS41", 23.1, 0.34, 0.55, 3.46, 2, "Yes", "Vesicle trafficking, lysosome"],
     ["ASAP3", 25.0, 0.32, 0.53, 1.60, 1, "Yes", "Arf GAP, cell migration"]],
    notes=["*K = 0.00 for SLC5A7, STRA8, TCEA3, and STIP1 is a degenerate boundary estimate (omega-saturation artifact),",
           "not a literal estimate of zero selection intensity. The gene-driven (relaxed) class requires FDR-significant 0 < K < 1,",
           "so these genes retain their dual-driven classification rather than being counted as relaxed."])

out2 = f"{BASE}/results/paper/supplementary/Supplementary_Text_Tables_v9.xlsx"
wb2.save(out2)
print("saved", out2)
