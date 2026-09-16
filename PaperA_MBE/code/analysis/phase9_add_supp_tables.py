# -*- coding: utf-8 -*-
"""Append sheets S4-4 (GO-slim composition) and S4-5 (component decomposition)
to Supplementary_Text_Tables_v9.xlsx (both copies: results/paper/supplementary
and package 07_supplementary_tables)."""
import shutil

import pandas as pd
from openpyxl import load_workbook

BASE = 'd:/人类正选择基因项目/'
SRC = BASE + 'results/paper/supplementary/Supplementary_Text_Tables_v9.xlsx'
PKG = BASE + 'MBE_submission_PaperA/07_supplementary_tables/Supplementary_Text_Tables_v9.xlsx'
OUT = BASE + '.workbuddy/tmp_supp_report.txt'
report = []

p2 = pd.read_csv(BASE + 'results/phase9_hardening/p2_goslim_composition.csv')
p3 = pd.read_csv(BASE + 'results/phase9_hardening/p3_component_decomposition.csv')

# P3 pivot: dominant-component counts per class + EA split summary
import json
with open(BASE + 'results/phase9_hardening/p3_component_decomposition.json', encoding='utf-8') as f:
    p3j = json.load(f)
rows = []
for cls in ('regulation_driven', 'gene_driven', 'dual_driven'):
    d = p3j[cls]
    for comp, n in d['dominant_component_counts'].items():
        rows.append({'class': cls, 'dominant_component': comp, 'count': n,
                     'n_class': d['n'], 'pct': round(100 * n / d['n'], 1)})
for grp in ('rd_ea_gwas', 'rd_non_ea'):
    d = p3j[grp]
    for comp, n in d['dominant_component_counts'].items():
        rows.append({'class': grp, 'dominant_component': comp, 'count': n,
                     'n_class': d['n'], 'pct': round(100 * n / d['n'], 1)})
p3_tab = pd.DataFrame(rows)

for target in (SRC, PKG):
    wb = load_workbook(target)
    # remove if re-running
    for sn in ('S4-4 GO-slim', 'S4-5 Components'):
        if sn in wb.sheetnames:
            del wb[sn]
    ws1 = wb.create_sheet('S4-4 GO-slim')
    ws1.append(['Table S4-4. GO-slim-style functional category composition '
                '(keyword buckets over GO BP 2023; Fisher vs rest of universe; BH across 40 tests)'])
    ws1.append(list(p2.columns))
    for _, r in p2.iterrows():
        ws1.append(list(r.values))
    ws2 = wb.create_sheet('S4-5 Components')
    ws2.append(['Table S4-5. Dominant contributing component of GDS/RDS composite scores '
                '(weight x component value; argmax), by class and by EA-GWAS overlap'])
    ws2.append(list(p3_tab.columns))
    for _, r in p3_tab.iterrows():
        ws2.append(list(r.values))
    wb.save(target)
    report.append(f'updated {target} (sheets: {wb.sheetnames})')

# also copy CSVs into package supplementary tables dir
for f in ('p2_goslim_composition.csv', 'p3_component_decomposition.csv'):
    shutil.copy2(BASE + 'results/phase9_hardening/' + f,
                 BASE + 'MBE_submission_PaperA/07_supplementary_tables/TableS4-4_45_source_' + f)
    report.append('copied ' + f)

with open(OUT, 'w', encoding='utf-8') as f:
    f.write('\n'.join(report))
